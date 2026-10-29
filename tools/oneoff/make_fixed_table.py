# -*- coding: utf-8 -*-
"""Generate a corrected VisionMaster variable table from AutoShop's compiled
symbol map (VarAddrMap.am).

Only variables whose current address is known are emitted, so importing the
result updates those rows by name and leaves everything else untouched.
"""
import io
import os
import struct
import collections

import openpyxl
from openpyxl import Workbook

XLSX_IN = r"C:\Users\16958\Desktop\Plc_变量表.xlsx"
AM = r"C:\Users\16958\Desktop\歌尔现场\小车\PLC\VarAddrMap.am"
CSV = r"C:\Users\16958\Desktop\歌尔现场\小车\PLC\VarNameAndAddr.csv"
OUT_DIR = r"D:\C#\VM\artifacts\render\20261029"
XLSX_OUT = os.path.join(OUT_DIR, "Plc_变量表_修正地址.xlsx")
REPORT = r"D:\C#\VM\artifacts\logs\20261029\fix_table_report.txt"

os.makedirs(OUT_DIR, exist_ok=True)
log = io.open(REPORT, "w", encoding="utf-8")


def w(s=""):
    log.write(s + "\n")


# ---- VarAddrMap.am ----
raw = open(AM, "rb").read()
pos = 30
am = {}
entries = []
while pos + 2 <= len(raw):
    n = struct.unpack_from("<H", raw, pos)[0]
    if n < 2 or pos + 2 + n + 4 > len(raw):
        break
    nm = raw[pos + 2:pos + 2 + n - 1]
    if raw[pos + 2 + n - 1] != 0:
        break
    addr = struct.unpack_from("<I", raw, pos + 2 + n)[0]
    entries.append((nm.decode("utf-8", "replace"), addr))
    pos += 2 + n + 4

for n, a in entries:
    am.setdefault(n.upper(), a)
w("VarAddrMap.am symbols: %d (unique %d), consumed %d/%d bytes" % (len(entries), len(am), pos, len(raw)))

# ---- compiled base map (for sanity: sizes in bits) ----
craw = open(CSV, "rb").read()
ctext = craw.decode("utf-16") if craw[:2] in (b"\xff\xfe", b"\xfe\xff") else craw.decode("utf-8", "replace")
sizes = {}
for line in ctext.splitlines():
    p = [x.strip() for x in line.split("\t")]
    if len(p) < 4 or not p[1].upper().startswith("0X"):
        continue
    if p[3].isdigit():
        sizes[p[0].upper()] = int(p[3])

# ---- old table ----
wb = openpyxl.load_workbook(XLSX_IN, data_only=True)
ws = wb["变量表"]
rows = list(ws.iter_rows(values_only=True))
header = list(rows[0])
data = [list(r) for r in rows[1:] if r[0]]
w("old table rows: %d" % len(data))

kept = []
dropped = []
changed = 0
same = 0
for r in data:
    name = str(r[0])
    key = name.upper()
    if key not in am:
        dropped.append((name, str(r[3]), str(r[4]), str(r[2])))
        continue
    new = am[key]
    old = str(r[4])
    try:
        oldv = int(old, 16)
    except ValueError:
        oldv = -1
    if oldv != new:
        changed += 1
    else:
        same += 1
    row = list(r)
    row[4] = "%X" % new
    kept.append(row)

w("kept (address known): %d  changed=%d  unchanged=%d  dropped=%d"
  % (len(kept), changed, same, len(dropped)))

# dropped grouped by root
roots = collections.Counter()
for name, area, off, dtype in dropped:
    root = name.split("[")[0].split(".")[0]
    roots[root] += 1
w("\n-- dropped roots (top 30) --")
for r, c in roots.most_common(30):
    w("  %-30s %d" % (r, c))

# sanity: each kept row must fit inside its compiled base size
bad = []
for row in kept:
    key = str(row[0]).upper()
    base = key.split("[")[0].split(".")[0]
    if base in sizes and base in am:
        bits = sizes[base]
        span = am[key] - am[base]
        if span < 0 or span >= bits:
            bad.append((row[0], hex(am[key]), base, bits, span))
w("\n-- out-of-range rows: %d --" % len(bad))
for b in bad[:20]:
    w("  %s" % (b,))

# ---- write xlsx ----
out = Workbook()
ws2 = out.active
ws2.title = "变量表"
ws2.append(header)
for row in kept:
    ws2.append(row)
out.save(XLSX_OUT)
w("\nwritten: %s (%d rows)" % (XLSX_OUT, len(kept)))

# ---- spot checks ----
w("\n-- spot checks --")
for name in ("Alarm3[79]", "Alarm1[0]", "Alarm3_All_1", "PLCHeart", "CyCurrentPage", "VirBtn.EstopBtn"):
    for row in kept:
        if str(row[0]) == name:
            w("  %-18s new=0x%-9s (old=%s)" % (name, row[4],
              next((str(x[4]) for x in data if str(x[0]) == name), "?")))
            break
    else:
        w("  %-18s NOT in corrected table" % name)

log.close()
print("done")
