# -*- coding: utf-8 -*-
"""Final validation: read a sample of corrected addresses from the live PLC."""
import io
import struct
import socket

import openpyxl

XLSX = r"D:\C#\VM\artifacts\render\20261029\Plc_变量表_修正地址.xlsx"
OLD = r"C:\Users\16958\Desktop\Plc_变量表.xlsx"
OUT = r"D:\C#\VM\artifacts\logs\20261029\final_validate.txt"

READ_TEMPLATE = bytes.fromhex("16000103010000000100010000004000600010000000")

log = io.open(OUT, "w", encoding="utf-8")


def w(s=""):
    log.write(s + "\n")


def build_read(bit_addr, bits):
    f = bytearray(READ_TEMPLATE)
    f[14:18] = struct.pack("<i", bit_addr)
    f[18] = bits & 0xFF
    f[19] = (bits >> 8) & 0xFF
    f[20] = (bits >> 16) & 0xFF
    return bytes(f)


s = socket.create_connection(("127.0.0.1", 12939), timeout=5)
s.settimeout(10)


def recv_exact(n):
    buf = b""
    while len(buf) < n:
        c = s.recv(n - len(buf))
        if not c:
            raise IOError("closed")
        buf += c
    return buf


def read_bits(addr, count):
    frame = build_read(addr, count)
    s.sendall(struct.pack("<H", len(frame)) + frame[2:])
    head = recv_exact(2)
    total = struct.unpack("<H", head)[0]
    resp = head + recv_exact(total - 2)
    if resp[8] == 0x0F:
        raise IOError("PLC err %d" % struct.unpack("<H", resp[10:12])[0])
    data = resp[22:]
    return [(data[i // 8] >> (i % 8)) & 1 for i in range(count)], data


def read_scalar(addr, dtype):
    if "布尔" in dtype:
        return read_bits(addr, 1)[0][0]
    if "短整数" in dtype:
        d = read_bits(addr, 16)[1]
        return struct.unpack("<h", d[:2])[0]
    if "整数" in dtype:
        d = read_bits(addr, 32)[1]
        return struct.unpack("<i", d[:4])[0]
    if "单精度" in dtype:
        d = read_bits(addr, 32)[1]
        return round(struct.unpack("<f", d[:4])[0], 4)
    return None


new_wb = openpyxl.load_workbook(XLSX, data_only=True)
new_rows = {str(r[0]): (str(r[2]), str(r[4])) for r in list(new_wb["变量表"].iter_rows(values_only=True))[1:] if r[0]}

old_wb = openpyxl.load_workbook(OLD, data_only=True)
old_rows = {str(r[0]): str(r[4]) for r in list(old_wb["变量表"].iter_rows(values_only=True))[1:] if r[0]}

w("corrected table rows: %d" % len(new_rows))

w("\n%-22s %-9s %-10s %-9s %-10s" % ("name", "old", "old value", "new", "new value"))
for name in ("Alarm3[79]", "Alarm3[78]", "Alarm3[80]", "Alarm1[0]", "Alarm3_All_1",
             "Alarm3_Tag", "PLCHeart", "CyCurrentPage", "PLCVersion", "VirBtn.EstopBtn"):
    if name not in new_rows:
        w("%-22s (not in corrected table)" % name)
        continue
    dtype, new_off = new_rows[name]
    old_off = old_rows.get(name, "?")
    try:
        ov = read_scalar(int(old_off, 16), dtype)
    except Exception as ex:
        ov = "ERR"
    try:
        nv = read_scalar(int(new_off, 16), dtype)
    except Exception as ex:
        nv = "ERR"
    w("%-22s 0x%-8s %-10s 0x%-8s %-10s" % (name, old_off, ov, new_off, nv))

# whole Alarm3 array at the corrected base
bits, _ = read_bits(0x2101140, 100)
w("\ncorrected Alarm3[0..99] @0x2101140 -> set indices: %s" % [i for i, b in enumerate(bits) if b])

log.close()
print("done")
