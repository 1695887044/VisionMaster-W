# -*- coding: utf-8 -*-
"""Remap the old VisionMaster variable table onto AutoShop's compiled address map.

Input : Plc_变量表.xlsx (old, stale offsets) + VarNameAndAddr.csv (compiled truth)
Output: artifacts/render/<today>/Plc_变量表_编译地址.xlsx  (same 20 columns, fixed offsets)
        + a text report with per-base mapping and anomalies.
"""
import io
import re
import collections
import os
import sys

import openpyxl
from openpyxl import Workbook

XLSX_IN = r"C:\Users\16958\Desktop\Plc_变量表.xlsx"
CSV_IN = r"C:\Users\16958\Desktop\歌尔现场\小车\PLC\VarNameAndAddr.csv"
OUT_DIR = r"D:\C#\VM\artifacts\render\20261029"
XLSX_OUT = os.path.join(OUT_DIR, "Plc_变量表_编译地址.xlsx")
REPORT = r"D:\C#\VM\artifacts\logs\20261029\remap_report.txt"

os.makedirs(OUT_DIR, exist_ok=True)
log = io.open(REPORT, "w", encoding="utf-8")


def w(s=""):
    log.write(s + "\n")


# ---------------- compiled map ----------------
raw = open(CSV_IN, "rb").read()
text = raw.decode("utf-16") if raw[:2] in (b"\xff\xfe", b"\xfe\xff") else raw.decode("utf-8", "replace")

compiled = {}
for line in text.splitlines():
    p = [x.strip() for x in line.split("\t")]
    if len(p) < 2 or not p[1].upper().startswith("0X"):
        continue
    try:
        addr = int(p[1], 16)
    except ValueError:
        continue
    bits = 0
    if len(p) > 3 and p[3].isdigit():
        bits = int(p[3])
    compiled[p[0]] = dict(addr=addr, bits=bits, desc=p[2] if len(p) > 2 else "")

w("compiled entries: %d" % len(compiled))

# ---------------- old table ----------------
wb = openpyxl.load_workbook(XLSX_IN, data_only=True)
ws = wb["变量表"]
rows = list(ws.iter_rows(values_only=True))
header = list(rows[0])
data = [list(r) for r in rows[1:] if r[0]]
w("old rows: %d  header=%d cols" % (len(data), len(header)))

NAME_RE = re.compile(r"^(?P<base>[^\[]+)(\[(?P<idx>\d+)\])?(?P<rest>\..*)?$")


def split_name(name):
    m = NAME_RE.match(str(name).strip())
    if not m:
        return str(name).strip(), None, ""
    return m.group("base").strip(), (int(m.group("idx")) if m.group("idx") else None), (m.group("rest") or "")


# index old rows by base
by_base = collections.defaultdict(list)
for r in data:
    base, idx, rest = split_name(r[0])
    by_base[base].append((r, idx, rest))

# compiled lookup: exact, then case-insensitive
ci = {k.lower(): k for k in compiled}


def find_compiled(base):
    if base in compiled:
        return compiled[base], base
    k = ci.get(base.lower())
    if k:
        return compiled[k], k
    return None, None


unmapped_bases = []
mapped_bases = []
for base in sorted(by_base):
    entry, real = find_compiled(base)
    if entry is None:
        unmapped_bases.append(base)
    else:
        mapped_bases.append((base, real, entry))

w("\nold bases: %d | mapped to compiled: %d | unmapped: %d"
  % (len(by_base), len(mapped_bases), len(unmapped_bases)))
w("\n-- unmapped bases (kept at old address, needs manual review) --")
for b in unmapped_bases:
    sample = by_base[b][0][0]
    w("  %-32s area=%-3s old=0x%-8s n=%d" % (b, sample[3], str(sample[4]), len(by_base[b])))

# ---------------- remap ----------------
TYPE_BITS = {"布尔 (bool)": 1, "短整数 (short)": 16, "整数 (int)": 32,
             "单精度 (float)": 32, "小数 (double)": 64, "文本 (string)": 8,
             "短整数数组 (short[])": 16, "整数数组 (int[])": 32,
             "单精度数组 (float[])": 32, "布尔数组 (bool[])": 1}

out_rows = []
anomalies = []
stats = collections.Counter()

for r in data:
    name = r[0]
    base, idx, rest = split_name(name)
    entry, real = find_compiled(base)

    old_off = int(str(r[4]), 16)

    if entry is None:
        # keep old
        new_off = old_off
        stats["unmapped"] += 1
    else:
        old_base = min(int(str(x[0][4]), 16) for x in by_base[base])
        intra = old_off - old_base
        new_off = entry["addr"] + intra
        stats["mapped"] += 1

        elem_bits = TYPE_BITS.get(str(r[2]), None)
        size = (elem_bits or 1) * max(1, int(str(r[8]))) if elem_bits else 1
        if entry["bits"] and intra + size > entry["bits"]:
            anomalies.append((name, "越界", "intra=0x%X size=%d bits=%d" % (intra, size, entry["bits"])))
            stats["anomaly"] += 1
        if intra < 0:
            anomalies.append((name, "负偏移", "intra=0x%X" % intra))
            stats["anomaly"] += 1

    row = list(r)
    row[4] = ("%X" % new_off) if entry is not None else str(r[4])
    out_rows.append(row)

w("\nstats: %s" % dict(stats))
w("\n-- anomalies (first 40) --")
for a in anomalies[:40]:
    w("  %s" % (a,))

# ---------------- write xlsx ----------------
out = Workbook()
ws2 = out.active
ws2.title = "变量表"
ws2.append(header)
for row in out_rows:
    ws2.append(row)
out.save(XLSX_OUT)
w("\nwritten: %s (%d rows)" % (XLSX_OUT, len(out_rows)))

# ---------------- spot check ----------------
w("\n-- spot checks --")
for name in ("Alarm3[79]", "Alarm1[0]", "Alarm3_Tag", "PLCHeart", "CyCurrentPage", "Rail_CyCurrentPage"):
    for r in data:
        if str(r[0]) == name:
            base, idx, rest = split_name(name)
            entry, _ = find_compiled(base)
            w("  %-20s old=0x%-9s -> new=0x%-9s (compiled base=0x%X)"
              % (name, str(r[4]), ("%X" % (entry["addr"] + (int(str(r[4]), 16) - min(int(str(x[0][4]), 16) for x in by_base[base])))) if entry else "?", entry["addr"] if entry else 0))
            break

log.close()
print("done")
